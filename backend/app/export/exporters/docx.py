"""DocxExporter：Document AST → DOCX（python-docx）。

v1 为文本优先：标题/段落/行内强调与链接/列表/引用/表格/代码块/数学文本均可导出；
function_plot 与 mermaid 保留源码占位并记 warning。中文字体通过 Normal 样式挂载
w:eastAsia=宋体，保证 Word 打开时中文正常显示；bold/italic 由 Word 原生渲染。
"""

from __future__ import annotations

from io import BytesIO

from docx import Document as DocxDocument
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.opc.constants import RELATIONSHIP_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Mm, Pt, RGBColor

from app.contracts import ExportOptions
from app.export.document import Document, DocumentNode, ExportResult
from app.export.exporters._common import (
    MERMAID_WARNING,
    PLOT_PLACEHOLDER_WARNING,
    RAW_HTML_WARNING,
    format_meta_value,
    safe_url,
)

_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"

_HEADING_SIZES = {1: 20, 2: 16, 3: 14, 4: 12, 5: 11, 6: 10.5}


def _plain_text(children: list[DocumentNode]) -> str:
    """递归拼接行内节点的纯文本，供标题/链接文字等需要纯文本处使用。"""
    parts: list[str] = []
    for child in children:
        if child.type == "text":
            parts.append(child.text)
        elif child.children:
            parts.append(_plain_text(child.children))
        elif child.text:
            parts.append(child.text)
    return "".join(parts)


class DocxExporter:
    """实现 DocumentExporter：递归渲染 Document AST 为 DOCX 字节流。"""

    def render(self, document: Document, options: ExportOptions) -> ExportResult:
        """同步渲染；CPU 密集，调用方应放入线程执行，避免阻塞事件循环。"""
        self._doc = DocxDocument()
        self._configure_normal_style()
        self._configure_page(options)
        warnings: list[str] = []

        self._render_header(document, options, warnings)
        self._render_children(document.children, warnings)

        buf = BytesIO()
        self._doc.save(buf)
        return ExportResult(content=buf.getvalue(), mime_type=_MIME, warnings=warnings)

    async def export(self, document: Document, options: ExportOptions) -> ExportResult:
        """契约要求的 async 接口；渲染本身同步，直接转发到 render。"""
        return self.render(document, options)

    def _configure_normal_style(self) -> None:
        """Normal 样式挂载 CJK 字体；拉丁用 Calibri，中文用宋体。"""
        style = self._doc.styles["Normal"]
        style.font.name = "Calibri"
        style.font.size = Pt(11)
        rfonts = style.element.get_or_add_rPr().get_or_add_rFonts()
        rfonts.set(qn("w:eastAsia"), "宋体")

    def _configure_page(self, options: ExportOptions) -> None:
        section = self._doc.sections[0]
        size = (options.page_size or "A4").lower()
        if size == "a4":
            section.page_width = Mm(210)
            section.page_height = Mm(297)
        elif size == "letter":
            section.page_width = Inches(8.5)
            section.page_height = Inches(11)

    # --- 文档头部 ---
    def _render_header(self, document: Document, options: ExportOptions, warnings: list[str]) -> None:
        title = str(document.attributes.get("title") or "")
        if options.include_title and title:
            p = self._doc.add_paragraph()
            run = p.add_run(title)
            run.bold = True
            run.font.size = Pt(22)
            p.paragraph_format.space_after = Pt(12)
        if options.include_metadata:
            metadata = document.attributes.get("metadata")
            if metadata:
                for key, value in metadata.items():
                    p = self._doc.add_paragraph()
                    run = p.add_run(f"{key}: {format_meta_value(value)}")
                    run.font.size = Pt(9)
                    run.font.color.rgb = RGBColor(0x57, 0x60, 0x6A)

    # --- 块级 ---
    def _render_children(self, children: list[DocumentNode], warnings: list[str]) -> None:
        for child in children:
            self._render_block(child, warnings)

    def _render_block(self, node: DocumentNode, warnings: list[str]) -> None:
        handler = getattr(self, f"_block_{node.type}", None)
        if handler is not None:
            handler(node, warnings)
        else:
            warnings.append(f"无法表示的节点类型已跳过：{node.type}")

    def _block_heading(self, node: DocumentNode, warnings: list[str]) -> None:
        level = max(1, min(6, int(node.attributes.get("level", 1))))
        p = self._doc.add_paragraph()
        run = p.add_run(_plain_text(node.children))
        run.bold = True
        run.font.size = Pt(_HEADING_SIZES[level])
        p.paragraph_format.space_before = Pt(14 if level <= 2 else 10)
        p.paragraph_format.space_after = Pt(6)

    def _block_paragraph(self, node: DocumentNode, warnings: list[str]) -> None:
        p = self._doc.add_paragraph()
        self._render_inline(p, node.children, warnings)

    def _block_blockquote(self, node: DocumentNode, warnings: list[str]) -> None:
        p = self._doc.add_paragraph()
        self._render_inline(p, node.children, warnings)
        p.paragraph_format.left_indent = Pt(16)
        for run in p.runs:
            run.font.color.rgb = RGBColor(0x57, 0x60, 0x6A)

    def _block_list(self, node: DocumentNode, warnings: list[str], level: int = 0) -> None:
        ordered = bool(node.attributes.get("ordered"))
        for index, item in enumerate(node.children, start=1):
            self._block_list_item(item, warnings, ordered, index, level)

    def _block_list_item(
        self,
        item: DocumentNode,
        warnings: list[str],
        ordered: bool,
        index: int,
        level: int,
    ) -> None:
        if item.attributes.get("task"):
            marker = "☑ " if item.attributes.get("checked") else "☐ "
        else:
            marker = f"{index}. " if ordered else "• "
        indent = Pt(18 + 18 * level)
        first = True
        for child in item.children:
            if child.type == "list":
                self._block_list(child, warnings, level + 1)
                continue
            if child.type == "paragraph":
                p = self._doc.add_paragraph()
                p.paragraph_format.left_indent = indent
                if first:
                    self._add_run(p, marker)
                    first = False
                self._render_inline(p, child.children, warnings)
            elif child.children:
                # 直接行内子节点：拼进一个段落
                p = self._doc.add_paragraph()
                p.paragraph_format.left_indent = indent
                if first:
                    self._add_run(p, marker)
                    first = False
                self._render_inline(p, child.children, warnings)
            else:
                self._render_block(child, warnings)
                first = False

    def _block_table(self, node: DocumentNode, warnings: list[str]) -> None:
        rows = node.children
        ncols = max((len(r.children) for r in rows), default=0)
        if not rows or ncols == 0:
            return
        table = self._doc.add_table(rows=len(rows), cols=ncols)
        table.style = "Table Grid"
        for ri, row in enumerate(rows):
            head = bool(row.attributes.get("head"))
            for ci in range(ncols):
                cell = table.cell(ri, ci)
                p = cell.paragraphs[0]
                if ci < len(row.children):
                    self._render_inline(p, row.children[ci].children, warnings, bold=head)

    def _block_code_block(self, node: DocumentNode, warnings: list[str]) -> None:
        lines = node.text.split("\n")
        p = self._doc.add_paragraph()
        self._shade_paragraph(p)
        p.paragraph_format.left_indent = Pt(8)
        p.paragraph_format.right_indent = Pt(8)
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after = Pt(8)
        for i, line in enumerate(lines):
            run = p.add_run(line)
            run.font.name = "Consolas"
            run.font.size = Pt(10)
            if i < len(lines) - 1:
                run.add_break()

    def _block_thematic_break(self, node: DocumentNode, warnings: list[str]) -> None:
        p = self._doc.add_paragraph()
        pPr = p._p.get_or_add_pPr()
        pBdr = OxmlElement("w:pBdr")
        bottom = OxmlElement("w:bottom")
        bottom.set(qn("w:val"), "single")
        bottom.set(qn("w:sz"), "6")
        bottom.set(qn("w:space"), "1")
        bottom.set(qn("w:color"), "D0D7DE")
        pBdr.append(bottom)
        pPr.append(pBdr)

    def _block_mermaid(self, node: DocumentNode, warnings: list[str]) -> None:
        warnings.append(MERMAID_WARNING)
        self._block_code_block(node, warnings)

    def _block_function_plot(self, node: DocumentNode, warnings: list[str]) -> None:
        warnings.append(PLOT_PLACEHOLDER_WARNING)
        self._block_code_block(node, warnings)

    def _block_math_block(self, node: DocumentNode, warnings: list[str]) -> None:
        p = self._doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run(f"$${node.text}$$")

    def _block_html_block(self, node: DocumentNode, warnings: list[str]) -> None:
        # 原始 HTML 不可信，按纯文本保留正文
        warnings.append(RAW_HTML_WARNING)
        self._doc.add_paragraph(node.text)

    # --- 行内（写入 run） ---
    def _render_inline(
        self,
        paragraph,
        children: list[DocumentNode],
        warnings: list[str],
        bold: bool = False,
        italic: bool = False,
    ) -> None:
        for child in children:
            self._render_inline_node(paragraph, child, warnings, bold, italic)

    def _render_inline_node(
        self, paragraph, node: DocumentNode, warnings: list[str], bold: bool, italic: bool
    ) -> None:
        t = node.type
        if t == "text":
            self._add_run(paragraph, node.text, bold=bold, italic=italic)
        elif t == "strong":
            self._render_inline(paragraph, node.children, warnings, bold=True, italic=italic)
        elif t == "emphasis":
            self._render_inline(paragraph, node.children, warnings, bold=bold, italic=True)
        elif t == "codespan":
            self._add_run(paragraph, node.text, code=True)
        elif t == "link":
            inner = _plain_text(node.children)
            href = str(node.attributes.get("href") or "")
            safe_href = safe_url(href)
            if safe_href is None:
                warnings.append(f"链接协议不安全，已降级为纯文本：{href!r}")
                self._render_inline(paragraph, node.children, warnings, bold, italic)
            else:
                self._add_hyperlink(paragraph, safe_href, inner)
        elif t == "image":
            src = str(node.attributes.get("src") or "")
            alt = str(node.attributes.get("alt") or "")
            if safe_url(src) is None:
                warnings.append(f"图片地址不安全，已跳过：{src!r}")
            else:
                warnings.append("图片未内嵌到 DOCX，已用替代文本表示")
            if alt:
                self._add_run(paragraph, alt)
        elif t == "math_inline":
            self._add_run(paragraph, f"\\({node.text}\\)")
        elif t == "linebreak":
            self._add_run(paragraph, "").add_break()
        else:
            warnings.append(f"无法表示的行内节点已跳过：{t}")

    def _add_run(self, paragraph, text: str, bold: bool = False, italic: bool = False, code: bool = False):
        run = paragraph.add_run(text)
        run.bold = bold
        run.italic = italic
        if code:
            run.font.name = "Consolas"
            run.font.size = Pt(10)
        return run

    def _add_hyperlink(self, paragraph, url: str, text: str) -> None:
        """写入可点击的超链接 run（python-docx 无公开 API，需手写 w:hyperlink）。"""
        part = paragraph.part
        r_id = part.relate_to(url, RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
        hyperlink = OxmlElement("w:hyperlink")
        hyperlink.set(qn("r:id"), r_id)
        run = OxmlElement("w:r")
        rPr = OxmlElement("w:rPr")
        rFonts = OxmlElement("w:rFonts")
        rFonts.set(qn("w:ascii"), "Calibri")
        rFonts.set(qn("w:hAnsi"), "Calibri")
        rFonts.set(qn("w:eastAsia"), "宋体")
        rPr.append(rFonts)
        color = OxmlElement("w:color")
        color.set(qn("w:val"), "0969DA")
        rPr.append(color)
        u = OxmlElement("w:u")
        u.set(qn("w:val"), "single")
        rPr.append(u)
        run.append(rPr)
        t = OxmlElement("w:t")
        t.text = text
        t.set(qn("xml:space"), "preserve")
        run.append(t)
        hyperlink.append(run)
        paragraph._p.append(hyperlink)

    def _shade_paragraph(self, paragraph, fill: str = "F2F2F2") -> None:
        """给段落加浅灰底纹，用于代码块占位。"""
        pPr = paragraph._p.get_or_add_pPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), fill)
        pPr.append(shd)
